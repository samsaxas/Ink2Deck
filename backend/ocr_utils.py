"""Gemini-only image text extraction and document export helpers.

No local OCR engine is used. Extraction tries configured Gemini models and
API keys, with bounded retries for transient failures.
"""
import os
import time
from io import BytesIO

from fpdf import FPDF
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")
GEMINI_PROMPT = (
    "Read this whiteboard or handwritten-notes image. Transcribe all visible "
    "text as accurately as possible, preserving headings, bullets, equations, "
    "symbols, line breaks, and the original language. Do not summarize or "
    "invent unreadable words. Mark genuinely unreadable portions as [unclear]. "
    "Return only the transcription."
)

_FONT_CANDIDATES = [
    os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf"),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]


class GeminiExtractionError(RuntimeError):
    """Raised when every configured Gemini attempt fails."""


def to_rgb(image):
    """Return an RGB PIL image."""
    return image.convert("RGB")


def _csv_env(name):
    return [part.strip() for part in os.getenv(name, "").split(",") if part.strip()]


def _api_keys(api_key=None):
    if api_key:
        return [api_key]
    keys = _csv_env("GEMINI_API_KEYS")
    primary = os.getenv("GEMINI_API_KEY", "").strip()
    if primary:
        keys.insert(0, primary)
    # Preserve order while removing duplicate values.
    return list(dict.fromkeys(keys))


def _models(model_name=None):
    primary = model_name or os.getenv("GEMINI_MODEL", GEMINI_MODEL)
    backups = _csv_env("GEMINI_FALLBACK_MODELS")
    return list(dict.fromkeys([primary, *backups]))


def extract_text_with_gemini(image, api_key=None, model_name=None):
    """Send the original image to Gemini; never invokes another OCR engine."""
    from google import genai
    from google.genai import types

    keys = _api_keys(api_key)
    if not keys:
        raise GeminiExtractionError(
            "Gemini API key is missing. Set GEMINI_API_KEY in the project .env file."
        )

    models = _models(model_name)
    timeout_ms = max(5, int(os.getenv("GEMINI_TIMEOUT_SECONDS", "60"))) * 1000
    retries = max(0, min(3, int(os.getenv("GEMINI_RETRIES", "1"))))
    retry_delay = max(0.0, float(os.getenv("GEMINI_RETRY_DELAY_SECONDS", "1.5")))
    buf = BytesIO()
    to_rgb(image).save(buf, format="PNG")
    image_part = types.Part.from_bytes(data=buf.getvalue(), mime_type="image/png")

    failures = []
    for model in models:
        for key_index, key in enumerate(keys):
            client = None
            for attempt in range(retries + 1):
                try:
                    client = genai.Client(
                        api_key=key,
                        http_options=types.HttpOptions(timeout=timeout_ms),
                    )
                    response = client.models.generate_content(
                        model=model,
                        contents=[GEMINI_PROMPT, image_part],
                    )
                    result = (getattr(response, "text", None) or "").strip()
                    if result:
                        return result
                    raise RuntimeError("Gemini returned an empty transcription.")
                except Exception as exc:
                    message = str(exc)
                    failures.append(f"{model} (key {key_index + 1}): {message[:240]}")
                    lowered = message.lower()
                    # Retry only transient/network/timeout/rate-limit/server failures.
                    transient = any(token in lowered for token in (
                        "timeout", "timed out", "temporarily", "connection",
                        "503", "502", "500", "429", "resource_exhausted",
                        "unavailable", "deadline",
                    ))
                    if attempt < retries and transient:
                        time.sleep(retry_delay * (2 ** attempt))
                        continue
                    break
                finally:
                    if client is not None:
                        try:
                            client.close()
                        except Exception:
                            pass

    details = " | ".join(failures[-6:])
    raise GeminiExtractionError(
        "Gemini could not extract this image after trying the configured "
        f"models and keys. Check API quota, model availability, and key validity. {details}"
    )


def extract_text(image, warn=None):
    """Return Gemini transcription only. Raises when Gemini attempts are exhausted."""
    text = extract_text_with_gemini(image)
    return text, "gemini"


def create_ppt_with_text_and_image(text, image):
    """Title slide, original image slide, then one slide per text paragraph."""
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Whiteboard Content"
    title_slide.placeholders[1].text = "Automatically Generated from Image"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    stream = BytesIO()
    to_rgb(image).save(stream, format="PNG")
    stream.seek(0)
    slide.shapes.add_picture(stream, Inches(1), Inches(1), width=Inches(6))

    for i, content in enumerate([s.strip() for s in text.split("\n\n") if s.strip()], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i}"
        slide.placeholders[1].text = content
        for paragraph in slide.placeholders[1].text_frame.paragraphs:
            paragraph.font.size = Pt(18)

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out


def create_pdf(text):
    """Return PDF bytes. Uses a Unicode font when available, else Latin-1 only."""
    pdf = FPDF()
    pdf.add_page()
    font_path = next((p for p in _FONT_CANDIDATES if os.path.exists(p)), None)
    if font_path:
        pdf.add_font("DejaVu", "", font_path)
        pdf.set_font("DejaVu", "", 12)
    else:
        pdf.set_font("Helvetica", "", 12)
        text = text.encode("latin-1", errors="replace").decode("latin-1")
    pdf.multi_cell(0, 10, text=text)
    return bytes(pdf.output())
